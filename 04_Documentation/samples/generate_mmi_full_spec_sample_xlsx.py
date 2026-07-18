#!/usr/bin/env python3
"""Generate the full-spec sample workbook for Markdown Matrix Injector docs.

Usage:
  python3 04_Documentation/samples/generate_mmi_full_spec_sample_xlsx.py
"""

from pathlib import Path

OUT_PATH = Path(__file__).with_name("mmi-full-spec-sample.xlsx")


def build_workbook() -> None:
    try:
        from openpyxl import Workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - explicit user guidance path
        raise SystemExit(
            "openpyxl is required. Install it with: python3 -m pip install --user openpyxl"
        ) from exc

    wb = Workbook()

    matrix = wb.active
    matrix.title = "Matrix"
    matrix.append(["ID", "Role", "Team", "Status", "Payload", "Owner", "Priority", "Tag"])
    rows = [
        ["REQ-001", "Admin", "Platform", "AUTH", '{"rule":"allow","scope":"all"}', "alice", "P1", "security"],
        ["REQ-002", "User", "Platform", "AUTH", '{"rule":"deny","scope":"guest"}', "bob", "P2", "security"],
        ["REQ-003", "Admin", "Data", "REVIEW", '{"sql":"select * from t"}', "carol", "P2", "query"],
        ["REQ-004", "Auditor", "Data", "AUTH", '{"rule":"allow","scope":"read"}', "dave", "P3", "audit"],
        ["REQ-005", "Admin", "Ops", "READY", '{"job":"nightly","retry":3}', "erin", "P1", "ops"],
        ["REQ-006", "User", "Ops", "BLOCKED", '{"job":"manual","reason":"approval"}', "frank", "P4", "ops"],
        ["REQ-007", "Admin", "Platform", "AUTH", '{"rule":"allow","scope":"regional"}', "grace", "P1", "security"],
        ["REQ-008", "Reviewer", "QA", "REVIEW", '{"case":"E2E","host":"linux"}', "heidi", "P2", "qa"],
        ["REQ-009", "Admin", "QA", "AUTH", '{"case":"unit","coverage":95}', "ivan", "P1", "qa"],
        ["REQ-010", "User", "Support", "READY", '{"ticket":"MMI-10"}', "judy", "P3", "support"],
        ["REQ-011", "Admin", "Support", "AUTH", '{"ticket":"MMI-11"}', "kate", "P2", "support"],
        ["REQ-012", "Architect", "Platform", "REVIEW", '{"decision":"ADR-0002"}', "leo", "P1", "architecture"],
    ]
    for row in rows:
        matrix.append(row)

    wide = wb.create_sheet("WideMatrix")
    wide.append([f"C{i:02d}" for i in range(1, 13)])  # 12 columns for wide-table fallback
    for row_index in range(1, 16):
        wide.append([f"R{row_index:02d}-V{col:02d}" for col in range(1, 13)])

    large = wb.create_sheet("LargeMatrix")
    large.append([f"H{col:02d}" for col in range(1, 26)])  # 25 columns
    for row_index in range(1, 26):
        large.append([f"L{row_index:02d}-{col:02d}" for col in range(1, 26)])  # 625 cells

    header_row = wb.create_sheet("HeaderRow")
    header_row.append(["ignore", "ignore", "ignore"])
    header_row.append(["Name", "Value", "Comment"])
    header_row.append(["FeatureA", "enabled", "phase2"])
    header_row.append(["FeatureB", "disabled", "phase3"])
    header_row.append(["FeatureC", "enabled", "phase4"])

    transpose = wb.create_sheet("Transpose")
    transpose.append(["Metric", "Q1", "Q2", "Q3", "Q4"])
    transpose.append(["Latency(ms)", 120, 110, 95, 90])
    transpose.append(["Errors", 5, 4, 2, 2])
    transpose.append(["Coverage", 80, 85, 90, 92])

    wb.save(OUT_PATH)


def main() -> None:
    build_workbook()
    print(f"generated: {OUT_PATH}")


if __name__ == "__main__":
    main()
